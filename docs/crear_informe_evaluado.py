"""Compone el informe APA con las métricas reales del experimento."""
from pathlib import Path
import json,re,shutil
from datetime import datetime,timezone
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1];EXP=ROOT/'experimentos/evaluacion_20261016'
WORK=Path.home()/'AppData/Local/Temp/informe_bovino_apa_20261016'
TITLE='Desarrollo y evaluación de una aplicación web para la clasificación de imágenes de piel y ojos bovinos'
OUT=ROOT/'docs/Informe_Bovino_APA_Evaluado.docx'
counts=json.loads((EXP/'conteos.json').read_text(encoding='utf-8'))
results_dir=sorted((EXP/'resultados').glob('*/resumen.json'))[-1].parent
results=json.loads((results_dir/'resumen.json').read_text(encoding='utf-8'))
train=results['resultados']['entrenamiento'];test=results['resultados']['prueba']
run=json.loads((EXP/'entrenamiento.json').read_text(encoding='utf-8'))
cfg=json.loads((EXP/'config_experimento.json').read_text(encoding='utf-8'))
labels=[c['etiqueta'] for c in cfg['clases']]
names=['Piel sana','Dermatosis nodular','Imagen no válida','Queratoconjuntivitis','Ojo sano']
def pct(x):return f'{100*x:.2f}'.replace('.',',')+'\u00a0%'
def number(x):return f'{x:,}'.replace(',','.')
correct=sum(test['matriz_confusion'][i][i] for i in range(5))
errors=[]
for i,row in enumerate(test['matriz_confusion']):
    for j,n in enumerate(row):
        if i!=j and n:errors.append((n,names[i].lower(),names[j].lower()))
errors.sort(reverse=True)
errtext='Las confusiones más frecuentes fueron: '+ '; '.join(f'{n} imágenes de {a} asignadas a {b}' for n,a,b in errors[:3])+'. La Figura 2 permite consultar los errores menos frecuentes. Las confusiones de una clase de enfermedad con su categoría saludable merecen atención al planificar la validación de campo.'
if not errors:errtext='Las 150 imágenes se clasificaron de acuerdo con sus etiquetas. Este resultado pertenece a la prueba interna y no permite asumir una ausencia de errores en nuevas condiciones.'
accepted=test['aceptadas_app'];acceptedacc=test['exactitud_aceptadas_app']
coverage=f'Además del resultado del modelo sobre toda la prueba, se aplicaron los controles de calidad y las reglas de aceptación. La aplicación aceptó {accepted} de las 150 imágenes, con una cobertura de {pct(test["cobertura_app"])}. La exactitud dentro de ese subconjunto fue {pct(acceptedacc) if acceptedacc is not None else "no calculable, porque no hubo imágenes aceptadas"}. La cobertura incluye las cinco clases en su denominador; las imágenes clasificadas como no válidas no se consideran respuestas aceptadas de una condición bovina.'
replacements={
    'EXACTITUD':pct(test['exactitud']),'PRECISION':pct(test['promedio_macro']['precision']),'SENSIBILIDAD':pct(test['promedio_macro']['recall']),
    'TOTAL_ENTRADAS':number(sum(c['entradas'] for c in counts.values())),
    'RESULTADOS_TEXTO':f'El modelo acertó {correct} de las 150 imágenes de prueba, lo que corresponde a una exactitud de {pct(test["exactitud"])}. La precisión macro fue {pct(test["promedio_macro"]["precision"])}, la sensibilidad macro {pct(test["promedio_macro"]["recall"])} y F1 macro {pct(test["promedio_macro"]["f1-score"])}. En las {number(train["cantidad"])} imágenes de entrenamiento, la exactitud fue {pct(train["exactitud"])}. La diferencia entre entrenamiento y prueba fue de {100*results["brecha_exactitud"]:.2f} puntos porcentuales. La Tabla 5 detalla el resultado de cada clase.'.replace('puntos porcentuales','puntos porcentuales'),
    'ERRORES_TEXTO':errtext,
    'COBERTURA_TEXTO':coverage,
    'DISCUSION_BRECHA':f'La exactitud de entrenamiento de {pct(train["exactitud"])} frente al {pct(test["exactitud"])} de prueba muestra una diferencia de {100*results["brecha_exactitud"]:.2f} puntos porcentuales. Esa brecha es compatible con sobreajuste o con diferencias entre los ejemplos de ambos conjuntos, aunque no identifica por sí sola la causa. El resultado justifica ampliar la validación y revisar la diversidad de las imágenes. Se conserva la medición obtenida, sin ajustar retrospectivamente el modelo para mejorar su respuesta sobre esta prueba.'.replace('14.73','14,73'),
}
replacements['RESULTADOS_TEXTO']=re.sub(r'(\d+)\.(\d+) puntos',r'\1,\2 puntos',replacements['RESULTADOS_TEXTO'])
sapi='https://universe.roboflow.com/muhammad-athar-althariq-irawan-m346dkx3922-wmdko/sapipise'
kamb='https://universe.roboflow.com/muhammad-athar-althariq-irawan-m346dkx3922-wmdko/kambingpise'
tables={
 'fases':('1','Aplicación de CRISP-DM en el proyecto',['Fase','Aplicación'],[1.75,4.75],[
 ['Comprensión del problema','Definición del apoyo visual y de las cinco categorías.'],
 ['Comprensión de los datos','Inventario, procedencia y revisión de especies y repeticiones.'],
 ['Preparación de los datos','Depuración, agrupación y reserva de 30 imágenes por clase.'],
 ['Modelado','Transferencia desde ImageNet y ajuste de cinco salidas.'],
 ['Evaluación','Predicciones, métricas por clase y matriz de confusión.'],
 ['Puesta en uso','Integración en una interfaz web con controles de entrada.']],
 'Elaboración propia basada en las fases de CRISP-DM descritas por IBM (s. f.).'),
 'procedencia':('2','Procedencia documentada de las imágenes recibidas',['Clase','Fuente','Enlace','Licencia','Cantidad'],[1.8,1.25,1.2,1.5,.75],[
 ['Piel sana','Mendeley y Kaggle', '[Mendeley](https://data.mendeley.com/datasets/w36hpf86j2/1)\nKaggle sin identificar','CC BY 4.0 solo para Mendeley','782'],
 ['Dermatosis nodular','Mendeley y Kaggle','[Mendeley](https://data.mendeley.com/datasets/w36hpf86j2/1)\nKaggle sin identificar','CC BY 4.0 solo para Mendeley','621'],
 ['Imagen no válida','Colección entregada','No documentado','No declarada en el archivo','648'],
 ['Queratoconjuntivitis','Roboflow Universe',f'[Sapipise]({sapi})\n[Kambingpise]({kamb})','Dominio público declarado','654'],
 ['Ojo sano','Roboflow Universe',f'[Sapipise]({sapi})\n[Kambingpise]({kamb})','Dominio público declarado','546']],
 'Fuentes de piel declaradas en la descripción inicial; no consta su distribución dentro de cada clase. La licencia de Kumar y Shastri (2022) corresponde solo a Mendeley. La licencia ocular procede de las exportaciones de Irawan (2023a, 2023b). Ojo rosado reúne 498 muestras de sapipise y 156 de kambingpise; ojo sano, 261 y 285. Las 441 muestras de kambingpise se excluyeron del experimento. Los enlaces identificados son seleccionables.'),
 'distribucion':('3','Inventario y distribución del conjunto experimental',['Clase','Entradas','Únicas','Entrenamiento','Prueba'],[1.8,1,1,1.7,1],[
 [n,str(counts[l]['entradas']),str(counts[l]['unicas_pixeles']),str(counts[l]['entrenamiento']),'30'] for n,l in zip(names,labels)]+[
 ['Total',str(sum(c['entradas'] for c in counts.values())),str(sum(c['unicas_pixeles'] for c in counts.values())),str(train['cantidad']),'150']],
 'Únicas indica imágenes distintas por sus píxeles dentro de cada clase, antes de excluir especies y conflictos. Las columnas de entrenamiento y prueba reflejan la depuración completa. Se excluyeron 18 entradas de grupos con etiquetas diferentes y las variantes adicionales de los grupos reservados. Las exclusiones se aplican de forma sucesiva, por lo que sus categorías no deben sumarse como recuentos independientes.'),
 'parametros':('4','Parámetros de entrenamiento y funcionamiento',['Parámetro','Valor'],[3.3,3.2],[
 ['Arquitectura y pesos iniciales','MobileNetV2, α = 0,35; ImageNet'],['Ajuste de la red','Base congelada; salida de cinco clases'],
 ['Parámetros totales y ajustables',f'{number(run["parametros_totales"])} y {number(run["parametros_entrenables"])}'],
 ['Épocas y tamaño de lote','40 y 32'],['Optimizador y tasa de aprendizaje','Adam; 0,001'],['Abandono y regularización L2','0,20 y 0,0001'],
 ['Tamaño y canales de entrada','224 × 224 píxeles; RGB'],['Normalización','División por 127,5 y resta de 1'],
 ['Semilla y reserva de prueba','42; 30 imágenes por clase'],['Umbral y margen de aceptación','0,70 y 0,15'],
 ['Límite del archivo y resolución total','12 MiB y 24 millones de píxeles'],['Brillo mínimo y umbral de textura','35 y 45']],
 'Elaboración propia a partir de la ejecución del entrenamiento y la configuración de la aplicación. El brillo utiliza una escala de 0 a 255. La textura se mide con la varianza del laplaciano. Estos umbrales operativos se fijaron sin optimizarlos con la prueba final.'),
 'metricas':('5','Resultados de clasificación por clase en la prueba reservada',['Clase','Precisión','Sensibilidad','F1','Imágenes'],[1.85,1.15,1.3,1.1,1.1],[
 [n,pct(test['por_clase'][l]['precision']),pct(test['por_clase'][l]['recall']),pct(test['por_clase'][l]['f1-score']),'30'] for n,l in zip(names,labels)]+[
 ['Promedio macro',pct(test['promedio_macro']['precision']),pct(test['promedio_macro']['recall']),pct(test['promedio_macro']['f1-score']),'150']],
 'Cálculos obtenidos con evaluar.py y scikit-learn. Cada clase dispone de 30 imágenes separadas antes del ajuste. El promedio macro da el mismo peso a las cinco clases. Las métricas corresponden a todas las imágenes de prueba, incluidas las que la interfaz rechazaría.'),
}

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=sec.footer_distance=Inches(.5)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12);normal.font.color.rgb=RGBColor(0,0,0)
fmt=normal.paragraph_format;fmt.line_spacing=2;fmt.space_before=fmt.space_after=Pt(0);fmt.first_line_indent=Inches(.5);fmt.widow_control=True
lang=OxmlElement('w:lang');lang.set(qn('w:val'),'es-BO');normal.element.get_or_add_rPr().append(lang)
for name in ['Title','Heading 1','Heading 2','Heading 3']:
    s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(12);s.font.bold=True;s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.first_line_indent=Inches(0);s.paragraph_format.line_spacing=2
    s.paragraph_format.space_before=s.paragraph_format.space_after=Pt(0);s.paragraph_format.keep_with_next=True
    s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER if name in ['Title','Heading 1'] else WD_ALIGN_PARAGRAPH.LEFT
hp=sec.header.paragraphs[0];hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT;hp.paragraph_format.first_line_indent=Inches(0)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');hp._p.append(field)
doc.core_properties.title=TITLE;doc.core_properties.author='Cristian Alejandro Carvajal Luna Orozco';doc.core_properties.subject='Informe de Inteligencia Artificial con evaluación interna'
doc.core_properties.keywords='bovinos; clasificación; inteligencia artificial; evaluación; APA'
doc.core_properties.comments='';doc.core_properties.modified=datetime.now(timezone.utc)
def text_into(p,text):
    # Inline italics and usable dataset links.
    for part in re.split(r'(\[[^\]]+\]\(https?://[^)]+\))',text):
        m=re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)',part)
        if m:
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(m[2],RT.HYPERLINK,is_external=True))
            r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'000000');pr.append(color)
            underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');pr.append(underline);r.append(pr)
            t=OxmlElement('w:t');t.text=m[1];r.append(t);link.append(r);p._p.append(link)
        else:
            for i,s in enumerate(part.split('|')):p.add_run(s).italic=bool(i%2)
def para(text='',style=None,indent=True):
    p=doc.add_paragraph(style=style)
    if not indent:p.paragraph_format.first_line_indent=Inches(0)
    text_into(p,text);return p
def caption(num,title,kind='Tabla'):
    p=para(f'{kind} {num}',indent=False);p.runs[0].bold=True;p.paragraph_format.keep_with_next=True
    p=para(title,indent=False);p.runs[0].italic=True;p.paragraph_format.keep_with_next=True
def note(text):
    p=para(indent=False);p.add_run('Nota. ').italic=True;text_into(p,text);p.paragraph_format.keep_together=True
def table(key):
    num,title,headers,widths,rows,nt=tables[key];caption(num,title)
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
    for edge in ['top','bottom','left','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single' if edge in ['top','bottom'] else 'nil');e.set(qn('w:sz'),'6');e.set(qn('w:color'),'000000');borders.append(e)
    pr.append(borders);margins=OxmlElement('w:tblCellMar')
    for edge,v in [('top',85),('bottom',85),('left',85),('right',85)]:
        e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(v));e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)
    for i,values in enumerate([headers]+rows):
        row=t.rows[0] if i==0 else t.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for j,(cell,value) in enumerate(zip(row.cells,values)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.keep_with_next=(key=='metricas' or i==0 or i>=len(rows)-1)
            if j>0 and key in ('distribucion','metricas'):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            text_into(p,value)
            for r in p.runs:r.font.size=Pt(11);r.bold=(i==0)
            if i==0:
                cb=OxmlElement('w:tcBorders');b=OxmlElement('w:bottom');b.set(qn('w:val'),'single');b.set(qn('w:sz'),'6');b.set(qn('w:color'),'000000');cb.append(b);cell._tc.get_or_add_tcPr().append(cb)
    note(nt)

for _ in range(3):para(indent=False)
para(TITLE,'Title',False);para(indent=False)
for line in ['Cristian Alejandro Carvajal Luna Orozco','Universidad Católica Boliviana','Sede Santa Cruz','Inteligencia Artificial','16 de octubre de 2026']:
    p=para(line,indent=False);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
text=(ROOT/'docs/informe_evaluado.md').read_text(encoding='utf-8')
for key,value in replacements.items():text=text.replace('{{'+key+'}}',value)
assert '{{' not in text
(WORK/'informe_evaluado_final.md').write_text(text,encoding='utf-8')
mode='body'
for line in text.splitlines():
    if not line.strip():continue
    if line=='@RESUMEN':doc.add_page_break();para('Resumen','Heading 1',False);mode='abstract'
    elif line.startswith('@PALABRAS '):
        p=para();p.add_run('Palabras clave: ').italic=True;p.add_run(line[10:])
    elif line=='@CUERPO':doc.add_page_break();para(TITLE,'Title',False);mode='body'
    elif line=='@REFERENCIAS':doc.add_page_break();para('Referencias','Heading 1',False);mode='refs'
    elif line.startswith('@TABLA '):table(line[7:])
    elif line.startswith('@FIGURA '):
        flow=line[8:]=='flujo'
        caption('1' if flow else '2','Flujo de procesamiento de una fotografía' if flow else 'Matriz de confusión de las 150 imágenes de prueba','Figura')
        p=para(indent=False);p.paragraph_format.keep_with_next=True
        path=WORK/'flujo.png' if flow else results_dir/'matriz_prueba.png'
        pic=p.add_run().add_picture(str(path),width=Inches(6.5 if flow else 5.65));p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        pic._inline.docPr.set('descr','Diagrama del flujo de la aplicación' if flow else 'Matriz de confusión con 30 imágenes por clase; valores de aciertos y errores medidos en la prueba reservada.')
        note('Elaboración propia. Los controles de calidad pueden solicitar otra fotografía antes de la inferencia.' if flow else 'Resultado generado por evaluar.py. C0: piel sana; C1: dermatosis nodular; C2: imagen no válida; C3: queratoconjuntivitis; C4: ojo sano. Las filas indican las clases reales y las columnas las predichas.')
    elif line.startswith('## '):doc.add_heading(line[3:],level=2)
    elif line.startswith('# '):doc.add_heading(line[2:],level=1)
    else:
        p=para(line,indent=mode!='abstract')
        if mode=='refs':p.paragraph_format.left_indent=Inches(.5);p.paragraph_format.first_line_indent=Inches(-.5);p.paragraph_format.keep_together=True

for root in [doc.styles.element,doc.element]:
    for el in root.iter(qn('w:rFonts')):
        for a in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme']:el.attrib.pop(qn('w:'+a),None)
    for el in list(root.iter(qn('w:pBdr'))):el.getparent().remove(el)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
    rf=doc.styles[name].element.get_or_add_rPr().get_or_add_rFonts()
    for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),'Times New Roman')
    for el in doc.styles[name].element.iter(qn('w:color')):el.attrib.pop(qn('w:themeColor'),None)
alltext=' '.join(t.text for t in doc.element.iter(qn('w:t')))
words=len(re.findall(r'\S+',alltext))
assert 5000<=words<=5500,f'Palabras fuera del rango: {words}'
doc.save(OUT);copy=Path.home()/'Downloads'/OUT.name;shutil.copyfile(OUT,copy)
(WORK/'verificacion_informe_evaluado.json').write_text(json.dumps({'palabras':words,'tablas':len(doc.tables),'figuras':len(doc.inline_shapes),'metricas':str(results_dir),'archivo':str(copy)},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'palabras':words,'tablas':len(doc.tables),'figuras':len(doc.inline_shapes),'archivo':str(copy)},ensure_ascii=False))
